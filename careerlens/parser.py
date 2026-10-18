import io
import zipfile
from pathlib import Path
from fastapi import HTTPException

MAX_FILE = 10 * 1024 * 1024


def parse_document(filename, content):
    suffix = Path(filename or '').suffix.lower()
    if suffix not in ('.txt', '.pdf', '.docx'):
        raise HTTPException(415, '支持 TXT、文本型 PDF 和 DOCX 文件')
    if not content or len(content) > MAX_FILE:
        raise HTTPException(413, '文件为空或超过 10 MB')
    try:
        if suffix == '.txt':
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = content.decode('gb18030')
        elif suffix == '.docx':
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError('DOCX 解压后超过 30 MB')
            doc = Document(io.BytesIO(content))
            text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])
        else:
            from pypdf import PdfReader
            pdf = PdfReader(io.BytesIO(content))
            if pdf.is_encrypted or len(pdf.pages) > 30:
                raise ValueError('不支持加密或超过 30 页的 PDF')
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
    except Exception as exc:
        raise HTTPException(422, f'文件解析失败：{str(exc)[:160]}。可复制正文后粘贴。') from exc
    text = text.replace('\x00', '').strip()
    if not text:
        raise HTTPException(422, '未提取到文字，扫描 PDF 请先使用 OCR，或粘贴正文')
    if len(text) > 15000:
        raise HTTPException(422, '提取文字超过 15000 字，请精简后导入')
    return {'text': text, 'filename': Path(filename).name, 'characters': len(text), 'warning': '请核对解析顺序与表格内容；保存简历限 5000 字，岗位限 3000 字。'}
