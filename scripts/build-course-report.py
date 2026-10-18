"""Build the editable course report from its maintained Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent.parent


def main():
    doc=Document()
    section=doc.sections[0]
    section.page_width=Inches(8.5)
    section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.85)
    section.left_margin=section.right_margin=Inches(.9)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number'):
        style=doc.styles[name]
        style.font.name='Calibri'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体')
        style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.line_spacing=1.25
    doc.styles['Normal'].font.size=Pt(11)
    for name,size in [('Title',23),('Subtitle',12),('Heading 1',16),('Heading 2',12.5)]:
        doc.styles[name].font.size=Pt(size)
    for name in ('Heading 1','Heading 2'):
        doc.styles[name].paragraph_format.space_before=Pt(15)
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.core_properties.title='基于本地大语言模型的求职材料诊断与模拟面试系统'
    doc.core_properties.subject='CareerLens 课程设计报告'
    doc.core_properties.author=''
    lines=(ROOT/'docs/课程报告.md').read_text(encoding='utf-8').splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:
            i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                parts=[p.strip() for p in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',p) for p in parts):rows.append(parts)
                i+=1
            table=doc.add_table(rows=len(rows),cols=len(rows[0]))
            table.alignment=WD_TABLE_ALIGNMENT.CENTER
            table.autofit=False
            widths=[1.2,2.0,3.5] if len(rows[0])==3 else [1.2,5.5]
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                edge_el=OxmlElement('w:'+edge)
                for k,v in [('val','single'),('sz','5'),('color','D9D9D9')]:edge_el.set(qn('w:'+k),v)
                borders.append(edge_el)
            table._tbl.tblPr.append(borders)
            for n,row in enumerate(rows):
                if n==0:
                    repeat=OxmlElement('w:tblHeader');table.rows[n]._tr.get_or_add_trPr().append(repeat)
                for j,value in enumerate(row):
                    cell=table.cell(n,j);cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr()
                    mar=OxmlElement('w:tcMar')
                    for edge in ('top','left','bottom','right'):
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
                    tcpr.append(mar)
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E9EEF2' if n==0 else 'FFFFFF');tcpr.append(fill)
                    cell.text=value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.15
                        for run in p.runs:run.font.size=Pt(10);run.bold=(n==0)
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
            continue
        if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
        elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
        elif line=='CareerLens 课程设计报告':doc.add_paragraph(line,'Subtitle')
        else:doc.add_paragraph(line)
        i+=1
    footer=section.footer.paragraphs[0]
    footer.alignment=2
    run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    doc.save(ROOT/'docs/课程报告.docx')
    print(ROOT/'docs/课程报告.docx')


if __name__=='__main__':main()
